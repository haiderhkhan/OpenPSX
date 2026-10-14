"""P2-05: validated upsert of a manual .csv / .xlsx file into data/openpsx.db (SCHEMA §2).

Usage: .venv\\Scripts\\python scripts/import_manual.py <kind> <file> [--dry-run]
Kinds: prices (kmi30_history = prices with symbol KMI30), fundamentals, macro, corporate_actions, news.
Templates live in data/manual/templates/ ('#' lines are comments; '# e.g.' lines show a valid row).

All-or-nothing: if any row is bad, every error is printed with its line number and nothing is written.
--dry-run runs the same transaction, prints the counts, then rolls back.
After a prices or corporate_actions import, inferred breaks, the bonus/split adjustment and the KMI30
back-fill are recomputed (P2-04, D-25, D-33).
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sqlite3
import sys
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))  # run as a script: make `backend` and `scripts` importable

from backend.config import DB_PATH, TZ  # noqa: E402
from backend.db import adjust  # noqa: E402
from backend.db.connection import connect_rw  # noqa: E402
from backend.ingest import news  # noqa: E402
from scripts import build_db  # noqa: E402

TEMPLATES = ROOT / "data" / "manual" / "templates"
ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
URL = re.compile(r"https?://\S+")
MANUAL = "manual"
REQUIRED = {
    "prices": ("symbol", "date", "close"),
    "fundamentals": ("symbol", "period_end", "period_type", "fiscal_year"),
    "macro": ("series", "date", "value"),
    "corporate_actions": ("symbol", "ex_date", "action_type", "value"),
    "news": ("published_at", "title", "url"),
}
KIND_ALIASES = {"kmi30_history": "prices"}
TABLES = {"prices": ("prices_daily", ("symbol", "date")),
          "fundamentals": ("fundamentals", ("symbol", "period_end", "period_type")),
          "macro": ("macro_series", ("series", "date")),
          "corporate_actions": ("corporate_actions", ("symbol", "ex_date", "action_type")),
          "news": ("news", ("url",))}
PERIOD_TYPES = {"Q1", "H1", "9M", "FY", "TTM"}
SERIES = set(build_db.MACRO_SERIES.values())
ACTIONS = {"CASH_DIVIDEND", "BONUS", "RIGHT", "SPLIT"}
NEWS_SOURCES = {"DAWN", "BR", "MANUAL"}


class RowError(ValueError):
    pass


# --- reading --------------------------------------------------------------------------------
def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_rows(path: Path) -> tuple[int, list[str], list[tuple[int, dict[str, str]]]]:
    """(header line, lower-case header, [(line number, {column: cell})]); blank and '#' lines skipped."""
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            raw = [(n, [_cell(v) for v in row])
                   for n, row in enumerate(book.worksheets[0].iter_rows(values_only=True), 1)]
        finally:
            book.close()
    elif suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:  # -sig: Excel writes a BOM
            reader = csv.reader(f)
            raw = [(reader.line_num, [c.strip() for c in row]) for row in reader]
    else:
        raise SystemExit(f"{path.name}: only .csv and .xlsx are supported")
    rows = [(n, cells) for n, cells in raw if any(cells) and not cells[0].startswith("#")]
    if not rows:
        raise SystemExit(f"{path.name}: no header row")
    (header_line, header), body = rows[0], rows[1:]
    keys = [h.lower() for h in header]
    return header_line, keys, [(n, dict(zip(keys, cells))) for n, cells in body]


# --- field parsers --------------------------------------------------------------------------
def _date(value: str, field: str, as_of: str | None = None) -> str:
    try:
        if not ISO_DATE.fullmatch(value):
            raise ValueError
        date.fromisoformat(value)
    except ValueError:
        raise RowError(f"{field} '{value}' is not a date (YYYY-MM-DD)") from None
    if as_of and value > as_of:
        raise RowError(f"{field} {value} is after AS_OF_DATE {as_of}")
    return value


def _num(value: str, field: str, required: bool = False, positive: bool = False) -> float | None:
    if value == "":
        if required:
            raise RowError(f"{field} is required")
        return None
    try:
        x = float(value.replace(",", ""))  # official downloads use thousand separators
    except ValueError:
        raise RowError(f"{field} '{value}' is not a number") from None
    if not math.isfinite(x) or (positive and x <= 0):
        raise RowError(f"{field} must be a positive number, got '{value}'")
    return x


def _choice(value: str, field: str, allowed: set[str]) -> str:
    if value.upper() not in allowed:
        raise RowError(f"{field} '{value}' is not one of {', '.join(sorted(allowed))}")
    return value.upper()


def _symbol(value: str, known: set[str]) -> str:
    if value.upper() not in known:
        raise RowError(f"unknown symbol '{value}'")
    return value.upper()


# --- row parsers: (count groups, record) ------------------------------------------------------
def parse_prices(r: dict[str, str], ctx: dict) -> tuple[list[str], dict]:
    symbol = _symbol(r.get("symbol", ""), ctx["symbols"])
    o, h, lo = (_num(r.get(k, ""), k, positive=True) for k in ("open", "high", "low"))
    close = _num(r.get("close", ""), "close", required=True, positive=True)
    if None not in (o, h, lo) and not (lo <= min(o, close) and max(o, close) <= h):
        raise RowError(f"OHLC inconsistent: needs low {lo} ≤ open {o}, close {close} ≤ high {h}")
    volume = _num(r.get("volume", ""), "volume")
    if volume is not None and (volume < 0 or not volume.is_integer()):
        raise RowError(f"volume must be a whole number ≥ 0, got '{r['volume']}'")
    return [symbol], {"symbol": symbol, "date": _date(r.get("date", ""), "date", ctx["as_of"]),
                      "open": o, "high": h, "low": lo, "close": close,
                      "volume": None if volume is None else int(volume), "adj_close": close, "source": MANUAL}


def parse_fundamentals(r: dict[str, str], ctx: dict) -> tuple[list[str], dict]:
    symbol = _symbol(r.get("symbol", ""), ctx["symbols"])
    year = r.get("fiscal_year", "")
    if not re.fullmatch(r"\d{4}", year):
        raise RowError(f"fiscal_year '{year}' is not a 4-digit year")
    return [symbol], {"symbol": symbol, "period_end": _date(r.get("period_end", ""), "period_end", ctx["as_of"]),
                      "period_type": _choice(r.get("period_type", ""), "period_type", PERIOD_TYPES),
                      "fiscal_year": int(year),
                      **{k: _num(r.get(k, ""), k) for k in ("eps", "revenue", "net_income", "dps")},
                      "source": r.get("source") or MANUAL}


def parse_macro(r: dict[str, str], ctx: dict) -> tuple[list[str], dict]:
    series = _choice(r.get("series", ""), "series", SERIES)
    observed = _date(r.get("date", ""), "date")  # macro may run past AS_OF: ML joins on release_date
    released = r.get("release_date", "")
    released = (_date(released, "release_date") if released
                else build_db.release_date(series, date.fromisoformat(observed)).isoformat())
    if released < observed:
        raise RowError(f"release_date {released} is before the observation date {observed}")
    return [series], {"series": series, "date": observed, "value": _num(r.get("value", ""), "value", required=True),
                      "release_date": released, "source": r.get("source") or MANUAL}


def parse_corporate_actions(r: dict[str, str], ctx: dict) -> tuple[list[str], dict]:
    symbol = _symbol(r.get("symbol", ""), ctx["symbols"])
    return [symbol], {"symbol": symbol, "ex_date": _date(r.get("ex_date", ""), "ex_date", ctx["as_of"]),
                      "action_type": _choice(r.get("action_type", ""), "action_type", ACTIONS),
                      "value": _num(r.get("value", ""), "value", required=True, positive=True),
                      "source": r.get("source") or MANUAL}


def parse_news(r: dict[str, str], ctx: dict) -> tuple[list[str], dict]:
    raw = r.get("published_at", "")
    try:
        if not ISO_DATE.match(raw):
            raise ValueError
        stamp = datetime.fromisoformat(raw)
    except ValueError:
        raise RowError(f"published_at '{raw}' is not YYYY-MM-DD or an ISO timestamp") from None
    stamp = stamp.replace(tzinfo=TZ) if stamp.tzinfo is None else stamp.astimezone(TZ)  # naive = PKT
    title, url = " ".join(r.get("title", "").split()), r.get("url", "")
    if not title:
        raise RowError("title is required")
    if not URL.fullmatch(url):
        raise RowError(f"url '{url}' is not an http(s) link")
    # The file's symbol column is authoritative; a blank one falls back to alias matching.
    given = [s for s in re.split(r"[|,;\s]+", r.get("symbol", "")) if s]
    symbols = sorted({_symbol(s, ctx["symbols"]) for s in given}) or news.match_symbols(title, ctx["matchers"])
    return symbols or ["(none)"], {
        "published_at": stamp.isoformat(timespec="seconds"),
        "source": _choice(r.get("source") or "MANUAL", "source", NEWS_SOURCES),
        "title": title, "url": url, "symbols": symbols, "fetched_at": ctx["now"]}


PARSERS = {"prices": parse_prices, "fundamentals": parse_fundamentals, "macro": parse_macro,
           "corporate_actions": parse_corporate_actions, "news": parse_news}


def validate(kind: str, header_line: int, header: list[str], rows: list[tuple[int, dict[str, str]]],
             ctx: dict) -> tuple[list[tuple[list[str], dict]], list[str]]:
    """Parsed records and 'line N: …' errors. A duplicate key is an error, except news: rows sharing a
    URL are merged into the first one (symbols unioned)."""
    missing = [c for c in REQUIRED[kind] if c not in header]
    if missing:
        return [], [f"line {header_line}: missing column(s) {', '.join(missing)}"]
    key_cols = TABLES[kind][1]
    records, errors, seen = [], [], {}
    for line, row in rows:
        try:
            groups, record = PARSERS[kind](row, ctx)
        except RowError as exc:
            errors.append(f"line {line}: {exc}")
            continue
        key = tuple(record[k] for k in key_cols)
        if key not in seen:
            seen[key] = line, (groups, record)
            records.append((groups, record))
        elif kind != "news":
            errors.append(f"line {line}: duplicate of line {seen[key][0]} ({', '.join(map(str, key))})")
        else:
            first_groups, first = seen[key][1]
            first["symbols"] = sorted(set(first["symbols"]) | set(record["symbols"]))
            first_groups[:] = first["symbols"] or ["(none)"]
    return records, errors


# --- writing --------------------------------------------------------------------------------
def upsert(conn: sqlite3.Connection, kind: str, records: list[tuple[list[str], dict]]) -> Counter:
    """Counter of (group, 'new' | 'replaced' | 'ignored'). Rows replace same-key rows (manual wins, D-25);
    news dedupes on URL: a known URL only gets the file's symbols, a known title + time is ignored."""
    table, key_cols = TABLES[kind]
    counts: Counter = Counter()
    for groups, record in records:
        row = {**record, "symbols": json.dumps(record["symbols"])} if kind == "news" else record
        exists = conn.execute(f"SELECT 1 FROM {table} WHERE " + " AND ".join(f"{k} = ?" for k in key_cols),
                              [row[k] for k in key_cols]).fetchone()
        cols, marks = ", ".join(row), ", ".join("?" * len(row))
        before = conn.total_changes
        if kind == "news":
            conn.execute(f"INSERT OR IGNORE INTO news ({cols}) VALUES ({marks}) "
                         "ON CONFLICT(url) DO UPDATE SET symbols = excluded.symbols", list(row.values()))
        else:
            conn.execute(f"INSERT OR REPLACE INTO {table} ({cols}) VALUES ({marks})", list(row.values()))
        status = "replaced" if exists else ("new" if conn.total_changes > before else "ignored")
        counts.update((g, status) for g in groups)
    return counts


def readjust(conn: sqlite3.Connection) -> dict:
    """Re-run P2-04 + D-25 after prices or corporate actions change. Inferred breaks are re-derived, so an
    official BONUS/SPLIT on the same day replaces the inferred one instead of adjusting twice."""
    conn.execute("DELETE FROM corporate_actions WHERE source = ?", (adjust.INFERRED,))
    inferred = adjust.infer_breaks(conn)
    adjust.apply_adjustments(conn)
    # ponytail: back-fill covers only dates before the earliest KMI30 row, so kmi30_history must be one
    # contiguous run ending where the source starts (R-06); a stray early row drops the back-fill after it.
    first_real, backfilled = build_db.backfill_index(conn)
    if backfilled:
        conn.execute("INSERT OR REPLACE INTO meta VALUES ('kmi30_backfilled_before', ?)", (first_real,))
    else:
        conn.execute("DELETE FROM meta WHERE key = 'kmi30_backfilled_before'")
    return {"inferred": len(inferred), "violations": adjust.ex_date_violations(conn),
            "backfilled": backfilled, "first_real": first_real}


def import_file(kind: str, path: Path, db_path: Path = DB_PATH, dry_run: bool = False) -> dict:
    """{'errors': [...]} on any bad row (nothing written), else counts + re-adjust summary."""
    kind = KIND_ALIASES.get(kind, kind)
    header_line, header, rows = read_rows(path)
    conn = connect_rw(db_path)
    try:
        ctx = {"symbols": {s for (s,) in conn.execute("SELECT symbol FROM symbols")},
               "as_of": conn.execute("SELECT value FROM meta WHERE key = 'as_of_date'").fetchone()[0],
               "matchers": news.load_matchers(conn), "now": datetime.now(TZ).isoformat(timespec="seconds")}
        records, errors = validate(kind, header_line, header, rows, ctx)
        if errors:
            return {"errors": errors}
        counts = upsert(conn, kind, records)
        redo = readjust(conn) if kind in ("prices", "corporate_actions") and records else None
        if dry_run:
            conn.rollback()
        else:
            conn.commit()
        return {"errors": [], "counts": counts, "readjust": redo, "dry_run": dry_run, "as_of": ctx["as_of"]}
    finally:
        conn.close()


def report(kind: str, path: Path, result: dict) -> None:
    if result["errors"]:
        print(f"REJECTED {path.name} ({kind}): {len(result['errors'])} error(s), nothing written", file=sys.stderr)
        for e in result["errors"]:
            print(f"  {e}", file=sys.stderr)
        return
    counts = result["counts"]
    print(f"{'DRY RUN, nothing written: ' if result['dry_run'] else 'Imported '}{path.name} as {kind} "
          f"(AS_OF_DATE {result['as_of']})")
    print(f"  {'group':10} {'rows':>6} {'new':>6} {'replaced':>9} {'ignored':>8}")
    for g in sorted({g for g, _ in counts}):
        n = {s: counts[(g, s)] for s in ("new", "replaced", "ignored")}
        print(f"  {g:10} {sum(n.values()):>6} {n['new']:>6} {n['replaced']:>9} {n['ignored']:>8}")
    print(f"  {'total':10} {sum(counts.values()):>6}  (a news row counts once per symbol)")
    if redo := result["readjust"]:
        print(f"  re-adjusted: {redo['inferred']} inferred breaks; KMI30 back-fill {redo['backfilled']} rows "
              f"before {redo['first_real']}; ex-date violators: {redo['violations'] or 'none'}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("kind", choices=sorted([*REQUIRED, *KIND_ALIASES]))
    parser.add_argument("file", type=Path)
    parser.add_argument("--db", type=Path, default=DB_PATH)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    result = import_file(args.kind, args.file, args.db, args.dry_run)
    report(args.kind, args.file, result)
    return 1 if result["errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
